import io
import os
import re
import shutil
import tempfile
from pathlib import Path

import pdfplumber
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _pil_image_from_stream(stream):
    from PIL import Image
    attrs = stream.attrs
    f = attrs.get("Filter")
    if isinstance(f, (tuple, list)):
        filters = list(f)
    elif f:
        filters = [f]
    else:
        filters = []
    data = stream.get_data()
    if "DCTDecode" in filters:
        return Image.open(io.BytesIO(data)), "jpeg", data
    if "JPXDecode" in filters:
        return Image.open(io.BytesIO(data)), "jp2", data
    w = int(attrs.get("Width", 1))
    h = int(attrs.get("Height", 1))
    n = len(data)
    px = max(1, w * h)
    mode = None
    if n == px:
        mode = "L"
    elif n == px * 3:
        mode = "RGB"
    elif n == px * 4:
        try:
            from pdfminer.pdftypes import resolve1
            cs = resolve1(attrs.get("ColorSpace"))
            if str(cs) == "CMYK":
                mode = "CMYK"
            else:
                mode = "RGBA"
        except Exception:
            mode = "RGBA"
    if mode:
        return Image.frombytes(mode, (w, h), data), "png", None
    try:
        return Image.open(io.BytesIO(data)), "png", None
    except Exception:
        return None, None, None


def _walk_images(node):
    from pdfminer.layout import LTImage
    found = []
    stack = [node]
    while stack:
        current = stack.pop()
        if isinstance(current, LTImage):
            found.append(current)
        else:
            try:
                stack.extend(list(current))
            except (TypeError, ValueError):
                pass
    return found


def collect_page_images(path):
    images = []
    try:
        from pdfminer.high_level import extract_pages
        with open(path, "rb") as fp:
            for ltp in extract_pages(fp):
                page_ims = [(tuple(im.bbox), im.stream)
                            for im in _walk_images(ltp)]
                images.append(page_ims)
    except Exception:
        images = []
    return images


def place_page_images(page_imgs, page_width, page_height, img_dir, page_no, log):
    images = []
    for bbox, stream in (page_imgs or []):
        try:
            x0, y0, x1, y1 = bbox
            w_pt = x1 - x0
            h_pt = y1 - y0
            if w_pt < 5 or h_pt < 5:
                continue
            if w_pt > 0.9 * page_width and h_pt > 0.9 * page_height:
                continue
            im, ext, raw = _pil_image_from_stream(stream)
            if im is None:
                continue
            if ext == "jpeg":
                data = raw
            elif ext == "jp2":
                buf = io.BytesIO()
                im.save(buf, format="JPEG")
                data = buf.getvalue()
                ext = "jpeg"
            else:
                buf = io.BytesIO()
                im.save(buf, format="PNG")
                data = buf.getvalue()
                ext = "png"
            path = os.path.join(img_dir, "p%d_i%d.%s" % (page_no, len(images), ext))
            with open(path, "wb") as fh:
                fh.write(data)
            PILImage.open(path).verify()
        except Exception:
            continue
        top = page_height - y1
        images.append({
            "path": path,
            "w_pt": w_pt,
            "h_pt": h_pt,
            "top": top,
            "center_x": (x0 + x1) / 2,
            "page_width": page_width,
        })
    return images


ACCENT = "#2D8CF0"
ACCENT_DARK = "#1f6fd0"
BG = "#ffffff"
PANEL = "#f4f6fb"
TEXT = "#1a2030"
MUTED = "#5f6b7d"
BORDER = "#d9e0ec"
GRID_COLOR = "9AA6B5"

CHART_WIDTH = 70
LINE_PT = 14
MAX_SPACER = 6

FORBIDDEN = set('[]:*?/\\')


def clean_sheet_name(name):
    cleaned = "".join("_" if c in FORBIDDEN else c for c in name).strip()
    if not cleaned:
        cleaned = "Sheet"
    return cleaned[:31]


def to_num(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    candidate = text.replace(",", "")
    if re.fullmatch(r"-?\d+(?:\.\d+)?", candidate):
        try:
            return int(candidate)
        except ValueError:
            return float(text.replace(",", ""))
    return text


def clamp(value, lo, hi):
    return max(lo, min(hi, value))


def rows_from_words(page):
    words = page.extract_words(keep_blank_chars=False)
    if not words:
        return []

    ordered = sorted(words, key=lambda w: (round(w["top"], 1), w["x0"]))
    lines = []
    current = []
    tolerance = 2.0
    for word in ordered:
        if current and abs(word["top"] - current[0]["top"]) <= tolerance:
            current.append(word)
        else:
            if current:
                lines.append(current)
            current = [word]
    if current:
        lines.append(current)

    rows = []
    for line in lines:
        line.sort(key=lambda w: w["x0"])
        gaps = [line[i]["x0"] - line[i - 1]["x1"] for i in range(1, len(line))]
        if gaps:
            gaps_sorted = sorted(gaps)
            median = gaps_sorted[len(gaps_sorted) // 2]
            threshold = max(median * 2.5, 4.0)
        else:
            threshold = float("inf")

        cells = []
        segment = []
        for i, word in enumerate(line):
            if segment and i > 0:
                gap = word["x0"] - line[i - 1]["x1"]
                if gap > threshold:
                    cells.append(segment)
                    segment = []
            segment.append(word)
        if segment:
            cells.append(segment)

        top = min(w["top"] for w in line)
        bottom = max(w["bottom"] for w in line)
        rows.append({
            "top": top,
            "bottom": bottom,
            "cells": [{"text": " ".join(w["text"] for w in seg), "words": seg}
                      for seg in cells],
        })
    rows.sort(key=lambda r: r["top"])
    return rows


def _line_seg(line):
    try:
        return {"x0": line["x0"], "x1": line["x1"], "top": line["top"], "bottom": line["bottom"]}
    except (TypeError, KeyError):
        try:
            return {"x0": line.x0, "x1": line.x1, "top": line.top, "bottom": line.bottom}
        except AttributeError:
            return None


def _rect_dict(rect):
    try:
        return {
            "x0": rect["x0"], "x1": rect["x1"], "top": rect["top"], "bottom": rect["bottom"],
            "fill": rect.get("fill"), "color": rect.get("non_stroking_color"),
        }
    except (TypeError, KeyError):
        try:
            return {
                "x0": rect.x0, "x1": rect.x1, "top": rect.top, "bottom": rect.bottom,
                "fill": rect.fill, "color": rect.non_stroking_color,
            }
        except AttributeError:
            return None


def _overlap(a0, a1, b0, b1):
    return max(0.0, min(a1, b1) - max(a0, b0))


def tuple_to_hex(color):
    if color is None:
        return None
    try:
        vals = list(color)
    except TypeError:
        return None
    if len(vals) < 3:
        return None
    r = clamp(int(float(vals[0]) * 255), 0, 255)
    g = clamp(int(float(vals[1]) * 255), 0, 255)
    b = clamp(int(float(vals[2]) * 255), 0, 255)
    return "#%02X%02X%02X" % (r, g, b)


def _fill_at(geom, rects):
    x0, top, x1, bottom = geom
    cx = (x0 + x1) / 2
    cy = (top + bottom) / 2
    best = None
    best_area = None
    for r in rects:
        if not r or not r.get("fill"):
            continue
        if not (r.get("x0") is not None and r.get("x1") is not None and r.get("top") is not None and r.get("bottom") is not None):
            continue
        if r["x0"] <= cx <= r["x1"] and r["top"] <= cy <= r["bottom"]:
            area = (r["x1"] - r["x0"]) * (r["bottom"] - r["top"])
            if best_area is None or area < best_area:
                best = r.get("color")
                best_area = area
    return tuple_to_hex(best)


def _edges(geom, segments):
    x0, top, x1, bottom = geom
    edges = set()
    spanh = max(0.1, x1 - x0)
    spanv = max(0.1, bottom - top)
    for s in segments:
        if s is None:
            continue
        if abs(s["top"] - s["bottom"]) <= 0.9:
            y = s["top"]
            if abs(y - top) <= 0.9 and _overlap(s["x0"], s["x1"], x0, x1) >= 0.6 * spanh:
                edges.add("top")
            if abs(y - bottom) <= 0.9 and _overlap(s["x0"], s["x1"], x0, x1) >= 0.6 * spanh:
                edges.add("bottom")
        else:
            x = s["x0"]
            if abs(x - x0) <= 0.9 and _overlap(s["top"], s["bottom"], top, bottom) >= 0.6 * spanv:
                edges.add("left")
            if abs(x - x1) <= 0.9 and _overlap(s["top"], s["bottom"], top, bottom) >= 0.6 * spanv:
                edges.add("right")
    return edges


def _cell_style(geom, segments, rects, is_header, words):
    style = {"bold": False, "fill": None, "color": None, "wrap": True, "edges": None}
    if is_header:
        style["bold"] = True
        style["fill"] = ACCENT
        style["color"] = "FFFFFF"
    if geom:
        fill = _fill_at(geom, rects)
        if fill:
            style["fill"] = fill
            try:
                r = int(fill[1:3], 16)
                g = int(fill[3:5], 16)
                b = int(fill[5:7], 16)
                lum = 0.299 * r + 0.587 * g + 0.114 * b
                style["color"] = "FFFFFF" if lum < 127 else "000000"
            except ValueError:
                pass
        if segments:
            edges = _edges(geom, segments)
            if edges:
                style["edges"] = sorted(edges)
    for w in (words or []):
        fn = (w.get("fontname") or "") + " " + (w.get("fontname_alt") or "")
        if "Bold" in fn or "Black" in fn or "Heavy" in fn:
            style["bold"] = True
            break
    return style


def _page_objects_for_bbox(page, bbox, page_area):
    bx0, btop, bx1, bbottom = bbox
    segments = []
    rects = []
    for ln in page.lines:
        g = _line_seg(ln)
        if not g:
            continue
        cx = (g["x0"] + g["x1"]) / 2
        cy = (g["top"] + g["bottom"]) / 2
        if bx0 - 2 <= cx <= bx1 + 2 and btop - 2 <= cy <= bbottom + 2:
            segments.append(g)
    for r in page.rects:
        d = _rect_dict(r)
        if not d or d["x0"] is None:
            continue
        w = d["x1"] - d["x0"]
        h = d["bottom"] - d["top"]
        thin = w <= 1.5 or h <= 1.5
        if page_area and not thin and (w * h) > 0.6 * page_area:
            continue
        cx = (d["x0"] + d["x1"]) / 2
        cy = (d["top"] + d["bottom"]) / 2
        if not (bx0 - 2 <= cx <= bx1 + 2 and btop - 2 <= cy <= bbottom + 2):
            continue
        if thin:
            if h <= 1.5:
                segments.append({"x0": d["x0"], "x1": d["x1"], "top": d["top"] + h / 2,
                                 "bottom": d["top"] + h / 2})
            else:
                segments.append({"x0": d["x0"] + w / 2, "x1": d["x0"] + w / 2,
                                 "top": d["top"], "bottom": d["bottom"]})
        else:
            rects.append({"x0": d["x0"], "x1": d["x1"], "top": d["top"],
                          "bottom": d["bottom"], "fill": d["fill"], "color": d["color"]})
    return segments, rects


def _words_in_geom(words, geom):
    x0, top, x1, bottom = geom
    result = []
    for w in words:
        wx = (w["x0"] + w["x1"]) / 2
        wy = (w["top"] + w["bottom"]) / 2
        if x0 - 0.5 <= wx <= x1 + 0.5 and top - 0.5 <= wy <= bottom + 0.5:
            result.append(w)
    return result


def _page_units(page, tablelist, page_no, img_dir, log):
    units = []
    bboxes = [t.bbox for t in tablelist]
    page_area = page.width * page.height
    page_words = page.extract_words(keep_blank_chars=False)

    for ti, table in enumerate(tablelist):
        try:
            data = table.extract()
        except Exception:
            data = []
        if not data:
            continue
        ncols = max(len(r) for r in data)
        segments, rects = _page_objects_for_bbox(page, table.bbox, page_area)
        row_cells = table.rows
        for r, row_data in enumerate(data):
            if r >= len(row_cells):
                top = None
                bottom = None
                geoms = []
            else:
                geoms = []
                for c in range(ncols):
                    if c < len(row_cells[r].cells) and row_cells[r].cells[c] is not None:
                        geoms.append(row_cells[r].cells[c])
                    else:
                        geoms.append(None)
                tops = [g[1] for g in geoms if g]
                bot = [g[3] for g in geoms if g]
                if not tops:
                    continue
                top = min(tops)
                bottom = max(bot)
            cells = [row_data[c] if c < len(row_data) else None for c in range(ncols)]
            styles = []
            for c in range(ncols):
                geom = geoms[c] if c < len(geoms) else None
                words = _words_in_geom(page_words, geom) if geom else []
                styles.append(_cell_style(geom, segments, rects, r == 0, words))
            units.append({
                "kind": "trow", "top": top, "bottom": bottom,
                "cells": cells, "styles": styles, "ncols": ncols,
            })

    lines = rows_from_words(page)
    for line in lines:
        cy = (line["top"] + line["bottom"]) / 2
        inside = False
        for (bx0, btop, bx1, bbottom) in bboxes:
            if btop - 2 <= cy <= bbottom + 2:
                inside = True
                break
        if inside:
            continue
        cells = [c["text"] for c in line["cells"]]
        styles = []
        for c in line["cells"]:
            st = {"bold": False, "fill": None, "color": None, "wrap": True, "edges": None}
            for w in c["words"]:
                fn = (w.get("fontname") or "") + " " + (w.get("fontname_alt") or "")
                if "Bold" in fn or "Black" in fn or "Heavy" in fn:
                    st["bold"] = True
                    break
            styles.append(st)
        units.append({
            "kind": "text", "top": line["top"], "bottom": line["bottom"],
            "cells": cells, "styles": styles, "ncols": len(cells),
        })
    return units


def assemble_sheet(pdf, doc_imgs, include_images, img_dir, name, log):
    rows = []
    images = []
    cursor = 1
    total = len(pdf.pages)

    for page_no, page in enumerate(pdf.pages, 1):
        tablelist = []
        try:
            tablelist = page.find_tables()
        except Exception as exc:
            log("  table scan error on page %d: %s" % (page_no, exc))

        units = _page_units(page, tablelist, page_no, img_dir, log)
        if include_images and doc_imgs:
            page_imgs = doc_imgs[page_no - 1] if page_no - 1 < len(doc_imgs) else []
            for im in place_page_images(page_imgs, page.width, page.height,
                                        img_dir, page_no, log):
                units.append({
                    "kind": "img", "top": im["top"], "bottom": im["top"] + im["h_pt"],
                    "img": im, "ncols": 0,
                })
        units.sort(key=lambda u: u["top"])

        if page_no > 1:
            rows.append({"cells": ["Page %d" % page_no], "styles": [{"bold": True, "color": MUTED, "wrap": False, "edges": ["bottom"]}]})
            cursor += 1
            for _ in range(2):
                rows.append({"cells": [], "styles": []})
                cursor += 1

        if units:
            lead = clamp(int(units[0]["top"] / LINE_PT), 0, MAX_SPACER)
            for _ in range(lead):
                rows.append({"cells": [], "styles": []})
                cursor += 1

        prev_bottom = None
        for u in units:
            if u["kind"] == "img":
                u["img"]["anchor_row"] = cursor
                hrows = clamp(int(u["img"]["h_pt"] / LINE_PT), 1, 40)
                images.append(u["img"])
                for _ in range(hrows):
                    rows.append({"cells": [], "styles": []})
                    cursor += 1
                prev_bottom = u["top"] + u["img"]["h_pt"]
                continue
            if prev_bottom is not None:
                gap = u["top"] - prev_bottom
                if gap > 8:
                    spacers = clamp(int(gap / 12), 1, MAX_SPACER)
                    for _ in range(spacers):
                        rows.append({"cells": [], "styles": []})
                        cursor += 1
            rows.append({"cells": u["cells"], "styles": u["styles"]})
            cursor += 1
            prev_bottom = u["bottom"]
        if units:
            log("  page %d/%d -> %d units" % (page_no, total, len(units)))
        else:
            log("  page %d/%d -> no extractable content" % (page_no, total))

    maxcols = 0
    for r in rows:
        maxcols = max(maxcols, len(r["cells"]))
    if maxcols == 0:
        maxcols = 1

    for im in images:
        col = maxcols
        if im.get("page_width"):
            idx = int(round((im["center_x"] / im["page_width"]) * maxcols))
            col = clamp(idx + 1, 1, maxcols)
        im["col"] = col
        im["row"] = clamp(im.get("anchor_row", 1), 1, max(1, len(rows)))
        im["w_px"] = int(im["w_pt"] * 96 / 72)
        im["h_px"] = int(im["h_pt"] * 96 / 72)

    return {"name": name, "rows": rows, "images": images, "maxcols": maxcols}


def write_workbook(sheet, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = clean_sheet_name(sheet["name"])
    rows = sheet["rows"]
    maxcols = sheet["maxcols"]

    widths = [8] * maxcols
    for r in rows:
        for c, text in enumerate(r["cells"]):
            if text is None:
                continue
            longest = max((len(x) for x in str(text).split("\n")), default=0)
            widths[c] = clamp(longest + 2, widths[c], CHART_WIDTH)

    thin = Side(style="thin", color=GRID_COLOR)

    for ri, r in enumerate(rows, 1):
        for c, text in enumerate(r["cells"]):
            st = r["styles"][c] if c < len(r["styles"]) else {}
            if text is None:
                continue
            cell = ws.cell(row=ri, column=c + 1, value=to_num(text))
            fill_hex = st.get("fill")
            if fill_hex:
                cell.fill = PatternFill("solid", fgColor=fill_hex.lstrip("#").upper())
            color = st.get("color") or "1A2030"
            cell.font = Font(size=11, bold=bool(st.get("bold")),
                             color=color.lstrip("#").upper())
            cell.alignment = Alignment(vertical="top", wrap_text=bool(st.get("wrap", True)))
            edges = st.get("edges")
            if edges:
                left = right = top = bottom = None
                if "left" in edges:
                    left = thin
                if "right" in edges:
                    right = thin
                if "top" in edges:
                    top = thin
                if "bottom" in edges:
                    bottom = thin
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    for c in range(maxcols):
        ws.column_dimensions[get_column_letter(c + 1)].width = widths[c]

    for im in sheet["images"]:
        if not os.path.exists(im["path"]):
            continue
        try:
            xl = XLImage(im["path"])
            xl.width = im["w_px"]
            xl.height = im["h_px"]
            anchor = "%s%d" % (get_column_letter(min(im["col"], maxcols)), im["row"])
            ws.add_image(xl, anchor)
        except Exception:
            continue

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    wb.save(output_path)


def convert_one(pdf_path, out_dir, log, include_images=True):
    path = Path(pdf_path)
    if not str(path).lower().endswith(".pdf"):
        return None
    base = path.stem
    candidate = Path(out_dir) / f"{base}.xlsx"
    n = 1
    while candidate.exists():
        candidate = Path(out_dir) / f"{base} ({n}).xlsx"
        n += 1
    log("Extracting content%s from %s" % (" and pictures" if include_images else "", path.name))
    img_dir = None
    try:
        if include_images:
            img_dir = tempfile.mkdtemp(prefix="pdf2excel_img_")
        with pdfplumber.open(pdf_path) as pdf:
            doc_imgs = collect_page_images(str(path)) if include_images else None
            sheet = assemble_sheet(pdf, doc_imgs, include_images, img_dir, base, log)
        write_workbook(sheet, str(candidate))
    finally:
        if img_dir:
            shutil.rmtree(img_dir, ignore_errors=True)
    return candidate
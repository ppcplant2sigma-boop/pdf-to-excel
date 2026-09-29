import os
import re
import queue
import shutil
import tempfile
import threading
import tkinter as tk
from pathlib import Path
from tkinter import ttk, filedialog, messagebox, scrolledtext

import pdfplumber
import pymupdf
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ACCENT = "#2D8CF0"
ACCENT_DARK = "#1f6fd0"
BG = "#ffffff"
PANEL = "#f4f6fb"
TEXT = "#1a2030"
MUTED = "#5f6b7d"
BORDER = "#d9e0ec"
GRID_COLOR = "9AA6B5"

CHART_WIDTH = 70
LINE_PT = 14
MAX_SPACER = 6

FORBIDDEN = set('[]:*?/\\')


def clean_sheet_name(name):
    cleaned = "".join("_" if c in FORBIDDEN else c for c in name).strip()
    if not cleaned:
        cleaned = "Sheet"
    return cleaned[:31]


def to_num(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    candidate = text.replace(",", "")
    if re.fullmatch(r"-?\d+(?:\.\d+)?", candidate):
        try:
            return int(candidate)
        except ValueError:
            return float(text.replace(",", ""))
    return text


def clamp(value, lo, hi):
    return max(lo, min(hi, value))


def rows_from_words(page):
    words = page.extract_words(keep_blank_chars=False)
    if not words:
        return []

    ordered = sorted(words, key=lambda w: (round(w["top"], 1), w["x0"]))
    lines = []
    current = []
    tolerance = 2.0
    for word in ordered:
        if current and abs(word["top"] - current[0]["top"]) <= tolerance:
            current.append(word)
        else:
            if current:
                lines.append(current)
            current = [word]
    if current:
        lines.append(current)

    rows = []
    for line in lines:
        line.sort(key=lambda w: w["x0"])
        gaps = [line[i]["x0"] - line[i - 1]["x1"] for i in range(1, len(line))]
        if gaps:
            gaps_sorted = sorted(gaps)
            median = gaps_sorted[len(gaps_sorted) // 2]
            threshold = max(median * 2.5, 4.0)
        else:
            threshold = float("inf")

        cells = []
        segment = []
        for i, word in enumerate(line):
            if segment and i > 0:
                gap = word["x0"] - line[i - 1]["x1"]
                if gap > threshold:
                    cells.append(segment)
                    segment = []
            segment.append(word)
        if segment:
            cells.append(segment)

        top = min(w["top"] for w in line)
        bottom = max(w["bottom"] for w in line)
        rows.append({
            "top": top,
            "bottom": bottom,
            "cells": [{"text": " ".join(w["text"] for w in seg), "words": seg}
                      for seg in cells],
        })
    rows.sort(key=lambda r: r["top"])
    return rows


def _line_seg(line):
    try:
        return {"x0": line["x0"], "x1": line["x1"], "top": line["top"], "bottom": line["bottom"]}
    except (TypeError, KeyError):
        try:
            return {"x0": line.x0, "x1": line.x1, "top": line.top, "bottom": line.bottom}
        except AttributeError:
            return None


def _rect_dict(rect):
    try:
        return {
            "x0": rect["x0"], "x1": rect["x1"], "top": rect["top"], "bottom": rect["bottom"],
            "fill": rect.get("fill"), "color": rect.get("non_stroking_color"),
        }
    except (TypeError, KeyError):
        try:
            return {
                "x0": rect.x0, "x1": rect.x1, "top": rect.top, "bottom": rect.bottom,
                "fill": rect.fill, "color": rect.non_stroking_color,
            }
        except AttributeError:
            return None


def _overlap(a0, a1, b0, b1):
    return max(0.0, min(a1, b1) - max(a0, b0))


def tuple_to_hex(color):
    if color is None:
        return None
    try:
        vals = list(color)
    except TypeError:
        return None
    if len(vals) < 3:
        return None
    r = clamp(int(float(vals[0]) * 255), 0, 255)
    g = clamp(int(float(vals[1]) * 255), 0, 255)
    b = clamp(int(float(vals[2]) * 255), 0, 255)
    return "#%02X%02X%02X" % (r, g, b)


def _fill_at(geom, rects):
    x0, top, x1, bottom = geom
    cx = (x0 + x1) / 2
    cy = (top + bottom) / 2
    best = None
    best_area = None
    for r in rects:
        if not r or not r.get("fill"):
            continue
        if not (r.get("x0") is not None and r.get("x1") is not None and r.get("top") is not None and r.get("bottom") is not None):
            continue
        if r["x0"] <= cx <= r["x1"] and r["top"] <= cy <= r["bottom"]:
            area = (r["x1"] - r["x0"]) * (r["bottom"] - r["top"])
            if best_area is None or area < best_area:
                best = r.get("color")
                best_area = area
    return tuple_to_hex(best)


def _edges(geom, segments):
    x0, top, x1, bottom = geom
    edges = set()
    spanh = max(0.1, x1 - x0)
    spanv = max(0.1, bottom - top)
    for s in segments:
        if s is None:
            continue
        if abs(s["top"] - s["bottom"]) <= 0.9:
            y = s["top"]
            if abs(y - top) <= 0.9 and _overlap(s["x0"], s["x1"], x0, x1) >= 0.6 * spanh:
                edges.add("top")
            if abs(y - bottom) <= 0.9 and _overlap(s["x0"], s["x1"], x0, x1) >= 0.6 * spanh:
                edges.add("bottom")
        else:
            x = s["x0"]
            if abs(x - x0) <= 0.9 and _overlap(s["top"], s["bottom"], top, bottom) >= 0.6 * spanv:
                edges.add("left")
            if abs(x - x1) <= 0.9 and _overlap(s["top"], s["bottom"], top, bottom) >= 0.6 * spanv:
                edges.add("right")
    return edges


def _cell_style(geom, segments, rects, is_header, words):
    style = {"bold": False, "fill": None, "color": None, "wrap": True, "edges": None}
    if is_header:
        style["bold"] = True
        style["fill"] = ACCENT
        style["color"] = "FFFFFF"
    if geom:
        fill = _fill_at(geom, rects)
        if fill:
            style["fill"] = fill
            try:
                r = int(fill[1:3], 16)
                g = int(fill[3:5], 16)
                b = int(fill[5:7], 16)
                lum = 0.299 * r + 0.587 * g + 0.114 * b
                style["color"] = "FFFFFF" if lum < 127 else "000000"
            except ValueError:
                pass
        if segments:
            edges = _edges(geom, segments)
            if edges:
                style["edges"] = sorted(edges)
    for w in (words or []):
        fn = (w.get("fontname") or "") + " " + (w.get("fontname_alt") or "")
        if "Bold" in fn or "Black" in fn or "Heavy" in fn:
            style["bold"] = True
            break
    return style


def _page_objects_for_bbox(page, bbox, page_area):
    bx0, btop, bx1, bbottom = bbox
    segments = []
    rects = []
    for ln in page.lines:
        g = _line_seg(ln)
        if not g:
            continue
        cx = (g["x0"] + g["x1"]) / 2
        cy = (g["top"] + g["bottom"]) / 2
        if bx0 - 2 <= cx <= bx1 + 2 and btop - 2 <= cy <= bbottom + 2:
            segments.append(g)
    for r in page.rects:
        d = _rect_dict(r)
        if not d or d["x0"] is None:
            continue
        w = d["x1"] - d["x0"]
        h = d["bottom"] - d["top"]
        thin = w <= 1.5 or h <= 1.5
        if page_area and not thin and (w * h) > 0.6 * page_area:
            continue
        cx = (d["x0"] + d["x1"]) / 2
        cy = (d["top"] + d["bottom"]) / 2
        if not (bx0 - 2 <= cx <= bx1 + 2 and btop - 2 <= cy <= bbottom + 2):
            continue
        if thin:
            if h <= 1.5:
                segments.append({"x0": d["x0"], "x1": d["x1"], "top": d["top"] + h / 2,
                                 "bottom": d["top"] + h / 2})
            else:
                segments.append({"x0": d["x0"] + w / 2, "x1": d["x0"] + w / 2,
                                 "top": d["top"], "bottom": d["bottom"]})
        else:
            rects.append({"x0": d["x0"], "x1": d["x1"], "top": d["top"],
                          "bottom": d["bottom"], "fill": d["fill"], "color": d["color"]})
    return segments, rects


def _words_in_geom(words, geom):
    x0, top, x1, bottom = geom
    result = []
    for w in words:
        wx = (w["x0"] + w["x1"]) / 2
        wy = (w["top"] + w["bottom"]) / 2
        if x0 - 0.5 <= wx <= x1 + 0.5 and top - 0.5 <= wy <= bottom + 0.5:
            result.append(w)
    return result


def _page_units(page, tablelist, page_no, img_dir, log):
    units = []
    bboxes = [t.bbox for t in tablelist]
    page_area = page.width * page.height
    page_words = page.extract_words(keep_blank_chars=False)

    for ti, table in enumerate(tablelist):
        try:
            data = table.extract()
        except Exception:
            data = []
        if not data:
            continue
        ncols = max(len(r) for r in data)
        segments, rects = _page_objects_for_bbox(page, table.bbox, page_area)
        row_cells = table.rows
        for r, row_data in enumerate(data):
            if r >= len(row_cells):
                top = None
                bottom = None
                geoms = []
            else:
                geoms = []
                for c in range(ncols):
                    if c < len(row_cells[r].cells) and row_cells[r].cells[c] is not None:
                        geoms.append(row_cells[r].cells[c])
                    else:
                        geoms.append(None)
                tops = [g[1] for g in geoms if g]
                bot = [g[3] for g in geoms if g]
                if not tops:
                    continue
                top = min(tops)
                bottom = max(bot)
            cells = [row_data[c] if c < len(row_data) else None for c in range(ncols)]
            styles = []
            for c in range(ncols):
                geom = geoms[c] if c < len(geoms) else None
                words = _words_in_geom(page_words, geom) if geom else []
                styles.append(_cell_style(geom, segments, rects, r == 0, words))
            units.append({
                "kind": "trow", "top": top, "bottom": bottom,
                "cells": cells, "styles": styles, "ncols": ncols,
            })

    lines = rows_from_words(page)
    for line in lines:
        cy = (line["top"] + line["bottom"]) / 2
        cx = None
        inside = False
        for (bx0, btop, bx1, bbottom) in bboxes:
            if btop - 2 <= cy <= bbottom + 2:
                inside = True
                break
        if inside:
            continue
        cells = [c["text"] for c in line["cells"]]
        styles = []
        for c in line["cells"]:
            st = {"bold": False, "fill": None, "color": None, "wrap": True, "edges": None}
            for w in c["words"]:
                fn = (w.get("fontname") or "") + " " + (w.get("fontname_alt") or "")
                if "Bold" in fn or "Black" in fn or "Heavy" in fn:
                    st["bold"] = True
                    break
            styles.append(st)
        units.append({
            "kind": "text", "top": line["top"], "bottom": line["bottom"],
            "cells": cells, "styles": styles, "ncols": len(cells),
        })
    return units


def place_page_images(pdf, page_no, img_dir, log):
    images = []
    try:
        page = pdf[page_no - 1]
        pw = page.rect.width
        ph = page.rect.height
        infos = page.get_image_info(xrefs=True)
    except Exception as exc:
        log("  image scan error on page %d: %s" % (page_no, exc))
        return images
    for info in infos:
        bbox = info.get("bbox")
        xref = info.get("xref") or 0
        if not bbox or xref <= 0:
            continue
        x0, y0, x1, y1 = bbox
        w_pt = x1 - x0
        h_pt = y1 - y0
        if w_pt < 5 or h_pt < 5:
            continue
        if w_pt > 0.9 * pw and h_pt > 0.9 * ph:
            continue
        try:
            extracted = pdf.extract_image(xref)
            data = extracted.get("image")
            ext = (extracted.get("ext") or "png").lower()
        except Exception:
            continue
        if not data:
            continue
        if ext == "jpg":
            ext = "jpeg"
        if ext not in ("png", "jpeg", "gif", "bmp"):
            ext = "png"
        path = os.path.join(img_dir, "p%d_x%d.%s" % (page_no, xref, ext))
        try:
            with open(path, "wb") as fh:
                fh.write(data)
            PILImage.open(path).verify()
        except Exception:
            continue
        images.append({
            "path": path,
            "w_pt": w_pt,
            "h_pt": h_pt,
            "top": y0,
            "center_x": (x0 + x1) / 2,
            "page_width": pw,
        })
    return images


def assemble_sheet(pdf, fdoc, include_images, img_dir, name, log):
    rows = []
    images = []
    cursor = 1
    total = len(pdf.pages)

    for page_no, page in enumerate(pdf.pages, 1):
        tablelist = []
        try:
            tablelist = page.find_tables()
        except Exception as exc:
            log("  table scan error on page %d: %s" % (page_no, exc))

        units = _page_units(page, tablelist, page_no, img_dir, log)
        if include_images and fdoc is not None:
            page_imgs = place_page_images(fdoc, page_no, img_dir, log)
            for im in page_imgs:
                units.append({
                    "kind": "img", "top": im["top"], "bottom": im["top"] + im["h_pt"],
                    "img": im, "ncols": 0,
                })
        units.sort(key=lambda u: u["top"])

        if page_no > 1:
            rows.append({"cells": ["Page %d" % page_no], "styles": [{"bold": True, "color": MUTED, "wrap": False, "edges": ["bottom"]}]})
            cursor += 1
            for _ in range(2):
                rows.append({"cells": [], "styles": []})
                cursor += 1

        if units:
            lead = clamp(int(units[0]["top"] / LINE_PT), 0, MAX_SPACER)
            for _ in range(lead):
                rows.append({"cells": [], "styles": []})
                cursor += 1

        prev_bottom = None
        for u in units:
            if u["kind"] == "img":
                u["img"]["anchor_row"] = cursor
                hrows = clamp(int(u["img"]["h_pt"] / LINE_PT), 1, 40)
                images.append(u["img"])
                for _ in range(hrows):
                    rows.append({"cells": [], "styles": []})
                    cursor += 1
                prev_bottom = u["top"] + u["img"]["h_pt"]
                continue
            if prev_bottom is not None:
                gap = u["top"] - prev_bottom
                if gap > 8:
                    spacers = clamp(int(gap / 12), 1, MAX_SPACER)
                    for _ in range(spacers):
                        rows.append({"cells": [], "styles": []})
                        cursor += 1
            rows.append({"cells": u["cells"], "styles": u["styles"]})
            cursor += 1
            prev_bottom = u["bottom"]
        if units:
            log("  page %d/%d -> %d units" % (page_no, total, len(units)))
        else:
            log("  page %d/%d -> no extractable content" % (page_no, total))

    maxcols = 0
    for r in rows:
        maxcols = max(maxcols, len(r["cells"]))
    if maxcols == 0:
        maxcols = 1

    for im in images:
        col = maxcols
        if im.get("page_width"):
            idx = int(round((im["center_x"] / im["page_width"]) * maxcols))
            col = clamp(idx + 1, 1, maxcols)
        im["col"] = col
        im["row"] = clamp(im.get("anchor_row", 1), 1, max(1, len(rows)))
        im["w_px"] = int(im["w_pt"] * 96 / 72)
        im["h_px"] = int(im["h_pt"] * 96 / 72)

    return {"name": name, "rows": rows, "images": images, "maxcols": maxcols}


def write_workbook(sheet, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = clean_sheet_name(sheet["name"])
    rows = sheet["rows"]
    maxcols = sheet["maxcols"]

    widths = [8] * maxcols
    for r in rows:
        for c, text in enumerate(r["cells"]):
            if text is None:
                continue
            longest = max((len(x) for x in str(text).split("\n")), default=0)
            widths[c] = clamp(longest + 2, widths[c], CHART_WIDTH)

    thin = Side(style="thin", color=GRID_COLOR)

    for ri, r in enumerate(rows, 1):
        for c, text in enumerate(r["cells"]):
            st = r["styles"][c] if c < len(r["styles"]) else {}
            if text is None:
                continue
            cell = ws.cell(row=ri, column=c + 1, value=to_num(text))
            fill_hex = st.get("fill")
            if fill_hex:
                cell.fill = PatternFill("solid", fgColor=fill_hex.lstrip("#").upper())
            color = st.get("color") or "1A2030"
            cell.font = Font(size=11, bold=bool(st.get("bold")),
                             color=color.lstrip("#").upper())
            cell.alignment = Alignment(vertical="top", wrap_text=bool(st.get("wrap", True)))
            edges = st.get("edges")
            if edges:
                left = right = top = bottom = None
                if "left" in edges:
                    left = thin
                if "right" in edges:
                    right = thin
                if "top" in edges:
                    top = thin
                if "bottom" in edges:
                    bottom = thin
                cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    for c in range(maxcols):
        ws.column_dimensions[get_column_letter(c + 1)].width = widths[c]

    for im in sheet["images"]:
        if not os.path.exists(im["path"]):
            continue
        try:
            xl = XLImage(im["path"])
            xl.width = im["w_px"]
            xl.height = im["h_px"]
            anchor = "%s%d" % (get_column_letter(min(im["col"], maxcols)), im["row"])
            ws.add_image(xl, anchor)
        except Exception:
            continue

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    wb.save(output_path)


def convert_one(pdf_path, out_dir, log, include_images=True):
    path = Path(pdf_path)
    if not str(path).lower().endswith(".pdf"):
        return None
    base = path.stem
    candidate = Path(out_dir) / f"{base}.xlsx"
    n = 1
    while candidate.exists():
        candidate = Path(out_dir) / f"{base} ({n}).xlsx"
        n += 1
    log("Extracting content%s from %s" % (" and pictures" if include_images else "", path.name))
    img_dir = None
    try:
        if include_images:
            img_dir = tempfile.mkdtemp(prefix="pdf2excel_img_")
        with pdfplumber.open(pdf_path) as pdf:
            fdoc = pymupdf.open(pdf_path) if include_images else None
            try:
                sheet = assemble_sheet(pdf, fdoc, include_images, img_dir, base, log)
            finally:
                if fdoc is not None:
                    fdoc.close()
        write_workbook(sheet, str(candidate))
    finally:
        if img_dir:
            shutil.rmtree(img_dir, ignore_errors=True)
    return candidate


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PDF to Excel")
        self.geometry("760x600")
        self.minsize(680, 560)
        self.configure(bg=BG)
        _res = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
        _ico = os.path.join(_res, "assets_win", "icon.ico")
        try:
            if os.path.exists(_ico):
                self.iconbitmap(_ico)
        except Exception:
            pass
        self.files = []
        self.q = queue.Queue()
        self.converting = False

        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("TFrame", background=BG)
        style.configure("Panel.TFrame", background=PANEL)
        style.configure("TLabel", background=BG, foreground=TEXT, font=("Segoe UI", 10))
        style.configure("Panel.TLabel", background=PANEL, foreground=TEXT, font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background=BG, foreground=MUTED, font=("Segoe UI", 9))
        style.configure("Dev.TLabel", background=BG, foreground=ACCENT_DARK, font=("Segoe UI", 9, "bold"))
        style.configure("Header.TLabel", background=BG, foreground=TEXT, font=("Segoe UI", 18, "bold"))
        style.configure("TButton", font=("Segoe UI", 10), padding=6)
        style.configure("Accent.TButton", font=("Segoe UI", 11, "bold"),
                        background=ACCENT, foreground="white", padding=(14, 9), borderwidth=0)
        style.map("Accent.TButton", background=[("active", ACCENT_DARK), ("disabled", "#9fb9d9")])
        style.configure("TEntry", font=("Segoe UI", 10))
        style.configure("TLabelframe", background=BG, bordercolor=BORDER, relief="solid")
        style.configure("TLabelframe.Label", background=BG, foreground=MUTED, font=("Segoe UI", 9, "bold"))
        style.configure("TCheckbutton", background=BG, foreground=TEXT, font=("Segoe UI", 10))
        style.configure("TListbox", font=("Segoe UI", 10), background="white", foreground=TEXT,
                        borderwidth=1, relief="solid")
        style.configure("TProgressbar", troughcolor="#e8ecf4", background=ACCENT, thickness=14)
        style.configure("Log.TText", font=("Consolas", 9), background="#10151f", foreground="#b9d8ff",
                        borderwidth=0, relief="flat")

        self._build_header()
        self._build_files()
        self._build_output()
        self._build_actions()
        self._build_log()
        self.after(120, self._drain_queue)

    def _build_header(self):
        frame = ttk.Frame(self)
        frame.pack(fill="x", padx=22, pady=(18, 6))
        ttk.Label(frame, text="PDF to Excel Converter", style="Header.TLabel").pack(anchor="w")
        ttk.Label(frame, text="Brings all tables, text and pictures of a PDF onto one Excel sheet, keeping the original layout. 100% offline.",
                  style="Muted.TLabel").pack(anchor="w", pady=(2, 0))
        ttk.Label(frame, text="Developer: Er. Durgesh Pandey  \u2764",
                  style="Dev.TLabel").pack(anchor="w", pady=(4, 0))

    def _build_files(self):
        frame = ttk.LabelFrame(self, text="1.  CHOOSE PDF FILES", padding=12)
        frame.pack(fill="both", expand=True, padx=22, pady=8)
        btn_row = ttk.Frame(frame)
        btn_row.pack(fill="x", pady=(0, 8))
        ttk.Button(btn_row, text="Add PDF Files...", command=self._add_files).pack(side="left")
        ttk.Button(btn_row, text="Add Folder...", command=self._add_folder).pack(side="left", padx=6)
        ttk.Button(btn_row, text="Clear", command=self._clear_files).pack(side="left")
        ttk.Label(btn_row, text="You can select multiple files.", style="Muted.TLabel").pack(side="right")

        opt_row = ttk.Frame(frame)
        opt_row.pack(fill="x", pady=(0, 8))
        self.include_images = tk.BooleanVar(value=True)
        ttk.Checkbutton(opt_row, text="Keep pictures from the PDF in their place",
                        variable=self.include_images).pack(side="left")
        ttk.Label(opt_row, text="Add an image-extraction pass.", style="Muted.TLabel").pack(side="right")

        body = ttk.Frame(frame)
        body.pack(fill="both", expand=True)
        self.listbox = tk.Listbox(
            body, exportselection=False, activestyle="none",
            bg="white", fg=TEXT, selectbackground="#dcebfc", selectforeground=TEXT,
            highlightthickness=1, highlightbackground=BORDER, relief="flat", font=("Segoe UI", 10))
        scroll = ttk.Scrollbar(body, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=scroll.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

    def _build_output(self):
        frame = ttk.LabelFrame(self, text="2.  OUTPUT DESTINATION", padding=12)
        frame.pack(fill="x", padx=22, pady=8)
        row = ttk.Frame(frame)
        row.pack(fill="x")
        self.same_folder = tk.BooleanVar(value=True)
        self.same_folder_chk = ttk.Checkbutton(
            row, text="Save next to each PDF (same folder)", variable=self.same_folder,
            command=self._toggle_output_entry)
        self.same_folder_chk.pack(side="left")
        ttk.Button(row, text="Browse...", command=self._browse_output).pack(side="right")
        self.output_var = tk.StringVar()
        self.output_entry = ttk.Entry(row, textvariable=self.output_var, state="disabled")
        self.output_entry.pack(side="right", fill="x", expand=True, padx=(0, 8))

    def _build_actions(self):
        frame = ttk.Frame(self)
        frame.pack(fill="x", padx=22, pady=10)
        self.convert_btn = ttk.Button(frame, text="Convert PDFs  \u2192  Excel", style="Accent.TButton",
                                      command=self._start_convert)
        self.convert_btn.pack(fill="x")
        self.progress = ttk.Progressbar(frame, mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(10, 4))
        self.status_var = tk.StringVar(value="Ready. Add PDF files to begin.")
        ttk.Label(frame, textvariable=self.status_var, style="Muted.TLabel").pack(anchor="w")
        ttk.Label(frame, text="Made with \u2764 by Er. Durgesh Pandey",
                  style="Dev.TLabel").pack(anchor="e", pady=(4, 0))

    def _build_log(self):
        frame = ttk.LabelFrame(self, text="  CONVERSION LOG  ", padding=12)
        frame.pack(fill="both", expand=True, padx=22, pady=(8, 18))
        self.log_text = scrolledtext.ScrolledText(
            frame, height=7, state="disabled", wrap="word",
            bg="#10151f", fg="#b9d8ff", insertbackground="white",
            font=("Consolas", 9), borderwidth=0, relief="flat")
        self.log_text.pack(fill="both", expand=True)

    def _toggle_output_entry(self):
        state = tk.NORMAL if not self.same_folder.get() else tk.DISABLED
        self.output_entry.config(state=state)

    def _add_files(self):
        selected = filedialog.askopenfilenames(
            title="Select PDF files",
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")])
        added = 0
        for f in selected:
            if f not in self.files:
                self.files.append(f)
                self.listbox.insert(tk.END, f)
                added += 1
        if added:
            self.status_var.set(f"{len(self.files)} file(s) ready.")

    def _add_folder(self):
        folder = filedialog.askdirectory(title="Select a folder containing PDF files")
        if not folder:
            return
        pdfs = sorted([str(p) for p in Path(folder).rglob("*.pdf")])
        added = 0
        for f in pdfs:
            if f not in self.files:
                self.files.append(f)
                self.listbox.insert(tk.END, f)
                added += 1
        self.status_var.set(f"Added {added} PDF(s) from folder. Total: {len(self.files)}.")

    def _clear_files(self):
        self.files.clear()
        self.listbox.delete(0, tk.END)
        self.status_var.set("Ready. Add PDF files to begin.")

    def _browse_output(self):
        folder = filedialog.askdirectory(title="Select output folder")
        if folder:
            self.output_var.set(folder)

    def _target_dir(self):
        if self.same_folder.get():
            return None
        folder = self.output_var.get().strip()
        if not folder:
            return Path.home() / "Desktop"
        return folder

    def _start_convert(self):
        if self.converting:
            return
        if not self.files:
            messagebox.showwarning("No files", "Please add at least one PDF file first.")
            return
        self.converting = True
        self.convert_btn.config(state="disabled")
        self._log("== Started converting %d file(s) ==" % len(self.files))
        thread = threading.Thread(target=self._convert_worker, daemon=True)
        thread.start()

    def _convert_worker(self):
        total = len(self.files)
        done = 0
        ok = 0
        failed = []
        use_images = self.include_images.get()

        def log(msg):
            self.q.put(("log", msg))

        for pdf in self.files:
            done += 1
            self.q.put(("status", "Converting %d/%d: %s" % (done, total, Path(pdf).name)))
            self.q.put(("progress", int((done - 1) / total * 100)))
            try:
                target = self._target_dir()
                target = target if target is not None else Path(pdf).parent
                Path(target).mkdir(parents=True, exist_ok=True)
                out = convert_one(pdf, target, log, use_images)
                if out:
                    log("Saved: %s" % out)
                    ok += 1
                else:
                    failed.append(pdf)
            except Exception as exc:
                detail = str(exc)
                if "password" in detail.lower() or "encrypted" in detail.lower():
                    log("ERROR: %s is password-protected. Open it and remove the password first." % Path(pdf).name)
                else:
                    log("ERROR: %s -> %s" % (Path(pdf).name, detail))
                failed.append(pdf)
            self.q.put(("progress", int(done / total * 100)))

        self.q.put(("progress", 100))
        self.q.put(("finished", (ok, len(failed), failed)))

    def _drain_queue(self):
        try:
            while True:
                item = self.q.get_nowait()
                kind = item[0]
                if kind == "log":
                    self._log(item[1])
                elif kind == "progress":
                    self.progress.configure(value=item[1])
                elif kind == "status":
                    self.status_var.set(item[1])
                elif kind == "finished":
                    ok, failed_count, _ = item[1]
                    self.converting = False
                    self.convert_btn.config(state="normal")
                    self.progress.configure(value=100)
                    if failed_count:
                        self.status_var.set("Finished. %d converted, %d failed." % (ok, failed_count))
                        messagebox.showwarning("Finished with errors",
                                               "%d file(s) converted successfully.\n%d file(s) failed.\n\nSee the log for details."
                                               % (ok, failed_count))
                    else:
                        self.status_var.set("All %d file(s) converted successfully!" % ok)
                        messagebox.showinfo("Done", "All %d file(s) converted successfully!" % ok)
        except queue.Empty:
            pass
        self.after(120, self._drain_queue)

    def _log(self, msg):
        if self.log_text:
            self.log_text.configure(state="normal")
            self.log_text.insert(tk.END, msg + "\n")
            self.log_text.see(tk.END)
            self.log_text.configure(state="disabled")


if __name__ == "__main__":
    app = App()
    app.mainloop()